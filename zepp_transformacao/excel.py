"""Construção do livro Excel nativo (openpyxl): tabelas, fórmulas vivas e gráficos."""
from __future__ import annotations

import math
import os
import re
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines, DateAxis
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.legend import Legend
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.chart.data_source import NumFmt
from openpyxl.chart.text import Text
from openpyxl.chart.title import Title
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun, Font as DrawingFont
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .config import Config
from .leituras import Selecao
from .medidas import CABECALHOS as CAB_MANUAIS, Semana
from .resumos import METRICAS_RESUMO

FONTE = "Arial"
COR = {
    "titulo": "243447", "texto": "263238", "suave": "52616B", "linha": "D7DEE3", "fundo": "F4F7F9",
    "cabecalho": "E8EDF3", "azul": "2F6BFF", "azul_claro": "EAF0FF", "laranja": "C76D00",
    "laranja_claro": "FFF3D6", "verde": "2E7D5B", "verde_claro": "E5F5EC", "vermelho": "B9382E",
    "vermelho_claro": "FCE8E6", "amarelo": "FFF7D6", "cinza_ponto": "9AA7B0",
}
FOLHA_CORPO, FOLHA_MANUAL = "Medições Corporais", "Medidas Manuais"
CORPO_COLUNAS = [
    ("Data/hora", 19, "dd/mm/yyyy hh:mm"), ("Peso (kg)", 11, "0.00"), ("IMC", 9, "0.00"),
    ("Gordura (%)", 11, "0.0"), ("Água corporal (%)", 13, "0.0"), ("Massa óssea (kg)", 12, "0.00"),
    ("Massa muscular (kg)", 13, "0.00"), ("Proteína (%)", 11, "0.0"), ("Massa gorda (kg)", 12, "0.00"),
    ("Gordura visceral (índice)", 13, "0"), ("Metabolismo basal (kcal/dia)", 15, "#,##0"),
    ("Idade corporal (anos)", 12, "0"), ("Pontuação corporal", 12, "0"),
    ("Peso, média 7 dias (kg)", 14, "0.00"), ("Gordura, média 7 dias (%)", 14, "0.00"),
    ("Músculo, média 7 dias (kg)", 14, "0.00"), ("Dia", 11, "dd/mm/yyyy"),
    # Instante UTC (oculto): identifica a última pesagem mesmo quando duas têm a mesma hora local
    # (mudança de hora de outubro) e mesmo que a tabela seja reordenada no Excel.
    ("Instante UTC", 19, "dd/mm/yyyy hh:mm:ss"),
]
CAMPO_CORPO = {"Peso (kg)": "Weight", "IMC": "BMI", "Gordura (%)": "BodyFat", "Água corporal (%)": "BodyWater",
               "Massa óssea (kg)": "BoneMass", "Massa muscular (kg)": "MuscleMass", "Proteína (%)": "ProteinRatio",
               "Gordura visceral (índice)": "VisceralFat", "Metabolismo basal (kcal/dia)": "BasalMetabolism",
               "Idade corporal (anos)": "MetabolicAge", "Pontuação corporal": "BodyScore"}
COLUNA_CORPO = {nome: get_column_letter(i + 1) for i, (nome, _, _) in enumerate(CORPO_COLUNAS)}
COLUNA_CAMPO = {**{campo: COLUNA_CORPO[nome] for nome, campo in CAMPO_CORPO.items()},
                "massa_gorda": COLUNA_CORPO["Massa gorda (kg)"]}
# Amplitude mínima dos eixos Y (unidade da métrica): evita ampliar ruído de medição.
AMPLITUDE_MINIMA = {"peso": 6, "gordura": 4, "musculo": 3, "cintura": 10, "pescoco": 3}
# Painel desenhado para o ecrã do utilizador (janela com ~1858 px úteis, 1 carácter de largura ≈ 9 px,
# 1 ponto de altura ≈ 1,667 px). Um único espaçamento (~27 px) em todo o lado: margens laterais,
# espaços entre cartões (horizontais e verticais), antes dos gráficos e entre gráficos.
ESPACO_LARGURA = 3          # caracteres ≈ 27 px
ESPACO_ALTURA = 16.2        # pontos ≈ 27 px
CARTAO_LARGURA = 15.9       # 3 colunas por cartão: 4 cartões + 5 espaços ≈ 1851 px
LINHA_CARTAO = 17           # pontos
# Linhas do Painel: 3/9/15/21 espaços entre blocos; gráficos ocupam 18 linhas e têm um espaço depois.
GRAFICOS_LINHAS = [(22, 39), (41, 58), (60, 77), (79, 96)]

borda = Side(style="thin", color=COR["linha"])


def eixo_legivel(valores, amplitude_minima: float) -> tuple[float, float, float]:
    """Mínimo, máximo e intervalo 'redondos' com margem de 12% e cerca de 5–7 marcas."""
    valores = [v for v in valores if isinstance(v, (int, float)) and math.isfinite(v)]
    if not valores:
        return 0, amplitude_minima, amplitude_minima / 5
    lo, hi = min(valores), max(valores)
    amplitude = max(hi - lo, amplitude_minima)
    baixo = max(0, (lo + hi - amplitude) / 2 - amplitude * 0.12)
    alto = (lo + hi + amplitude) / 2 + amplitude * 0.12
    expoente = math.floor(math.log10((alto - baixo) / 5))
    candidatos = []
    for potencia in range(expoente - 1, expoente + 2):
        for fator in (1, 2, 2.5, 5, 10):
            passo = fator * 10 ** potencia
            fundo, topo = math.floor(baixo / passo) * passo, math.ceil(alto / passo) * passo
            marcas = round((topo - fundo) / passo) + 1
            candidatos.append((abs(marcas - 6) + (0 if 5 <= marcas <= 7 else 10), topo - fundo, fundo, topo, passo))
    _, _, fundo, topo, passo = min(candidatos)
    return round(fundo, 10), round(topo, 10), round(passo, 10)


# ---------- estilos ----------

def _titulo(ws, texto: str, ultima_coluna: int, subtitulo: str | None = None, coluna: str = "A") -> None:
    """coluna = onde começam título e subtítulo ("B" quando a coluna A está fixa, para a linha
    de fixação do Excel não atravessar o texto)."""
    ws.sheet_view.showGridLines = False
    fim = get_column_letter(ultima_coluna)
    ws.merge_cells(f"{coluna}1:{fim}1")
    ws[f"{coluna}1"] = texto
    ws[f"{coluna}1"].font = Font(name=FONTE, bold=True, size=16, color=COR["titulo"])
    ws[f"{coluna}1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    if subtitulo:
        ws.merge_cells(f"{coluna}2:{fim}2")
        ws[f"{coluna}2"] = subtitulo
        ws[f"{coluna}2"].font = Font(name=FONTE, size=9, italic=True, color=COR["suave"])
        ws[f"{coluna}2"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 20


def _cabecalho(ws, linha: int, nomes: list[str], grupos: dict[str, str] | None = None) -> None:
    for i, nome in enumerate(nomes, start=1):
        c = ws.cell(linha, i, nome)
        c.font = Font(name=FONTE, bold=True, size=10, color=COR["texto"])
        c.fill = PatternFill("solid", fgColor=(grupos or {}).get(nome, COR["cabecalho"]))
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=borda)
    ws.row_dimensions[linha].height = 58


def _tabela(ws, nome: str, ref: str) -> None:
    t = Table(displayName=nome, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)


def _corpo_texto(ws, inicio_linha: int, fim_linha: int, colunas: int) -> None:
    for linha in ws.iter_rows(min_row=inicio_linha, max_row=fim_linha, max_col=colunas):
        for c in linha:
            c.font = Font(name=FONTE, size=10, color=COR["texto"])
            c.alignment = Alignment(vertical="center", horizontal=c.alignment.horizontal,
                                    wrap_text=c.alignment.wrap_text)


def _larguras(ws, larguras: list[float]) -> None:
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura


def _nota(ws, celula: str, fim: str, texto: str) -> None:
    ws.merge_cells(f"{celula}:{fim}")
    ws[celula] = texto
    ws[celula].font = Font(name=FONTE, size=9, italic=True, color=COR["suave"])
    ws[celula].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[ws[celula].row].height = 30


def _data_excel(expr: str) -> str:
    # TEXT() depende do idioma do Excel (aaaa vs yyyy); DAY/MONTH/YEAR não.
    return f'RIGHT("0"&DAY({expr}),2)&"/"&RIGHT("0"&MONTH({expr}),2)&"/"&YEAR({expr})'


# ---------- gráficos ----------

def _texto_grafico(tamanho: int = 900, negrito: bool = False, cor: str = COR["suave"]) -> RichText:
    props = CharacterProperties(sz=tamanho, b=negrito, solidFill=cor, latin=DrawingFont(typeface=FONTE))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=props), endParaRPr=props)])


def _titulo_grafico(texto: str, tamanho: int = 1200, negrito: bool = True, cor: str = COR["titulo"]) -> Title:
    props = CharacterProperties(sz=tamanho, b=negrito, solidFill=cor, latin=DrawingFont(typeface=FONTE))
    paragrafo = Paragraph(pPr=ParagraphProperties(defRPr=props), r=[RegularTextRun(rPr=props, t=texto)])
    return Title(tx=Text(rich=RichText(p=[paragrafo])), overlay=False)


def _estilo_grafico(ch, titulo: str, largura_cm: float, altura_cm: float) -> None:
    ch.title = _titulo_grafico(titulo)
    ch.width, ch.height = largura_cm, altura_cm
    ch.legend = None
    for eixo in (ch.x_axis, ch.y_axis):
        eixo.delete = False
        eixo.txPr = _texto_grafico()
        eixo.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=COR["linha"]))
    ch.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="EEF1F4")))
    ch.x_axis.majorGridlines = None


def _escala_y(ch, valores, amplitude: float, formato: str = "0.0") -> None:
    fundo, topo, passo = eixo_legivel(valores, amplitude)
    ch.y_axis.scaling.min, ch.y_axis.scaling.max, ch.y_axis.majorUnit = fundo, topo, passo
    ch.y_axis.numFmt = NumFmt(formatCode=formato, sourceLinked=False)


def _escala_datas(eixo, inicio: float, fim: float) -> None:
    eixo.numFmt = NumFmt(formatCode="dd/mm", sourceLinked=False)
    eixo.scaling.min, eixo.scaling.max = math.floor(inicio), math.ceil(fim)
    eixo.majorUnit = 28


def _linha(cor: str, largura_pt: float = 2.0, tracejado: str | None = None) -> GraphicalProperties:
    ln = LineProperties(solidFill=cor, w=int(largura_pt * 12700), prstDash=tracejado)
    return GraphicalProperties(ln=ln)


def grafico_diario(ws_corpo, ultima: int, coluna_valor: str, coluna_media: str, titulo: str, cor: str,
                   valores, medias, amplitude: float, rotulo: str) -> ScatterChart:
    """Leituras diárias em pontos e média móvel de 7 dias em linha."""
    ch = ScatterChart(scatterStyle="lineMarker")
    _estilo_grafico(ch, titulo, 30, 8.5)
    x = Reference(ws_corpo, min_col=1, min_row=4, max_row=ultima)
    pontos = Series(Reference(ws_corpo, range_string=f"'{FOLHA_CORPO}'!${coluna_valor}$4:${coluna_valor}${ultima}"),
                    x, title="Leituras")
    pontos.marker = Marker(symbol="circle", size=4)
    pontos.marker.graphicalProperties = GraphicalProperties(solidFill=COR["cinza_ponto"], ln=LineProperties(noFill=True))
    pontos.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    media = Series(Reference(ws_corpo, range_string=f"'{FOLHA_CORPO}'!${coluna_media}$4:${coluna_media}${ultima}"),
                   x, title="Média móvel de 7 dias")
    media.marker = Marker(symbol="none")
    media.graphicalProperties = _linha(cor, 2.25)
    media.smooth = False
    pontos.smooth = False
    ch.series += [pontos, media]
    ch.legend = Legend(legendPos="t", overlay=False, txPr=_texto_grafico())
    _escala_y(ch, list(valores) + list(medias), amplitude, "0.0")
    ch.y_axis.title = _titulo_grafico(rotulo, 900, False, COR["suave"])
    return ch


def grafico_manual(ws_manual, semanas: list[Semana], coluna: str, campo: str, titulo: str, cor: str,
                   amplitude: float) -> LineChart:
    """Linha contínua entre leituras consecutivas, tracejado sobre semanas sem medição."""
    ultima = len(semanas) + 3
    ch = LineChart()
    _estilo_grafico(ch, titulo, 15, 8)
    ch.add_data(Reference(ws_manual, range_string=f"'{FOLHA_MANUAL}'!${coluna}$4:${coluna}${ultima}"))
    ch.set_categories(Reference(ws_manual, range_string=f"'{FOLHA_MANUAL}'!$B$4:$B${ultima}"))
    serie = ch.series[0]
    serie.smooth = False
    serie.graphicalProperties = _linha(cor)
    serie.marker = Marker(symbol="circle", size=5)
    serie.marker.graphicalProperties = GraphicalProperties(solidFill=cor, ln=LineProperties(solidFill=cor))
    medidos = [i for i, s in enumerate(semanas) if getattr(s, campo) is not None]
    for anterior, atual in zip(medidos, medidos[1:]):
        if atual - anterior > 1:
            # O formato de um ponto aplica-se ao segmento que termina nele.
            ponto = DataPoint(idx=atual, marker=Marker(symbol="circle", size=5))
            ponto.marker.graphicalProperties = GraphicalProperties(solidFill=cor, ln=LineProperties(solidFill=cor))
            ponto.graphicalProperties = _linha(cor, 2.0, "dash")
            serie.dPt.append(ponto)
    ch.display_blanks = "span"
    ch.x_axis = DateAxis(crossAx=100, delete=False, txPr=_texto_grafico(),
                         spPr=GraphicalProperties(ln=LineProperties(solidFill=COR["linha"])))
    ch.x_axis.baseTimeUnit = "days"
    ch.x_axis.majorTimeUnit = "days"
    datas = [s.data for s in semanas if s.data is not None] or [s.inicio for s in semanas]
    _escala_datas(ch.x_axis, _serie(min(datas)), _serie(max(datas)))
    ch.y_axis.crossAx = 500
    _escala_y(ch, [getattr(s, campo) for s in semanas], amplitude, "0.0")
    return ch


def _serie(momento) -> float:
    base = datetime(1899, 12, 30)
    if not isinstance(momento, datetime):
        momento = datetime(momento.year, momento.month, momento.day)
    return (momento - base).total_seconds() / 86400


# ---------- livro ----------

def construir(caminho: Path, cfg: Config, selecao: Selecao, semanas: list[Semana], semanais: list[dict],
              mensais: list[dict], moveis: list[dict], gerado_em: datetime) -> None:
    leituras = selecao.aceites
    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.properties.creator = "zepp-life-excel-ia"
    wb.properties.title = "Transformação corporal"
    painel = wb.active
    painel.title = "Painel"
    mensal = wb.create_sheet("Resumo Mensal")
    semanal = wb.create_sheet("Resumo Semanal")
    manual = wb.create_sheet(FOLHA_MANUAL)
    corpo = wb.create_sheet(FOLHA_CORPO)
    criterios = wb.create_sheet("Critérios e Qualidade")
    inicio_txt = f"{cfg.inicio:%d/%m/%Y}"
    data_inicio = f"DATE({cfg.inicio.year},{cfg.inicio.month},{cfg.inicio.day})"
    atualizado = f"Atualizado em {gerado_em:%d/%m/%Y %H:%M}"

    # --- Medições corporais ---
    n = len(leituras)
    ultima = n + 3
    _titulo(corpo, f"Medições corporais completas desde {inicio_txt}", len(CORPO_COLUNAS),
            f"{n} leituras aceites · {atualizado} · fórmulas: massa gorda e médias móveis de 7 dias")
    _cabecalho(corpo, 3, [c for c, _, _ in CORPO_COLUNAS])
    rng = lambda col: f"'{FOLHA_CORPO}'!${COLUNA_CORPO[col]}$4:${COLUNA_CORPO[col]}${ultima}"
    A, Q, U = rng("Data/hora"), rng("Dia"), rng("Instante UTC")
    for i, l in enumerate(leituras):
        r = i + 4
        corpo.cell(r, 1, l.local)
        for nome, campo in CAMPO_CORPO.items():
            corpo[f"{COLUNA_CORPO[nome]}{r}"] = l[campo]
        corpo[f"I{r}"] = f'=IF(D{r}="","",B{r}*D{r}/100)'
        for nome, origem in (("Peso, média 7 dias (kg)", "B"), ("Gordura, média 7 dias (%)", "D"),
                             ("Músculo, média 7 dias (kg)", "G")):
            corpo[f"{COLUNA_CORPO[nome]}{r}"] = (
                f'=AVERAGEIFS(${origem}$4:${origem}${ultima},$A$4:$A${ultima},">"&$A{r}-7,$A$4:$A${ultima},"<="&$A{r})')
        corpo[f"Q{r}"] = f"=INT(A{r})"
        corpo[f"R{r}"] = l.instante_utc.replace(tzinfo=None)
    for i, (_, largura, formato) in enumerate(CORPO_COLUNAS, start=1):
        corpo.column_dimensions[get_column_letter(i)].width = largura
        for r in range(4, ultima + 1):
            corpo.cell(r, i).number_format = formato
    _corpo_texto(corpo, 4, ultima, len(CORPO_COLUNAS))
    _tabela(corpo, "tblCorpo", f"A3:R{ultima}")
    corpo.column_dimensions["Q"].hidden = True
    corpo.column_dimensions["R"].hidden = True
    corpo.freeze_panes = "A4"

    # --- Medidas manuais ---
    m_ultima = len(semanas) + 3
    _titulo(manual, "Medidas manuais semanais", 7,
            "Preencha as células amarelas (cintura ao nível do umbigo, pescoço). A data tem de pertencer à semana. "
            "Guarde e feche o Excel antes de atualizar.")
    _cabecalho(manual, 3, CAB_MANUAIS)
    for i, s in enumerate(semanas):
        r = i + 4
        manual[f"A{r}"] = s.rotulo
        manual[f"B{r}"] = s.data
        manual[f"C{r}"] = s.cintura
        manual[f"E{r}"] = s.pescoco
        manual[f"G{r}"] = s.notas
        manual[f"D{r}"] = (f'=IF(OR(C{r}="",COUNTIFS(B$4:B${m_ultima},"<"&B{r},C$4:C${m_ultima},">0")=0),"",'
                           f'C{r}-SUMIFS(C$4:C${m_ultima},B$4:B${m_ultima},'
                           f'_xlfn.MAXIFS(B$4:B${m_ultima},B$4:B${m_ultima},"<"&B{r},C$4:C${m_ultima},">0")))')
        manual[f"F{r}"] = (f'=IF(AND(C{r}="",E{r}=""),IF(G{r}<>"","Sem medição","Por preencher"),'
                           f'IF(AND(C{r}<>"",E{r}<>""),"Completa","Incompleta"))')
    _corpo_texto(manual, 4, m_ultima, 7)
    amarelo = PatternFill("solid", fgColor=COR["amarelo"])
    for r in range(4, m_ultima + 1):
        manual[f"B{r}"].number_format = "dd/mm/yyyy"
        manual[f"D{r}"].number_format = '+0.0;-0.0;0.0'
        for col in "CE":
            manual[f"{col}{r}"].number_format = "0.0"
        for col in "BCEG":
            manual[f"{col}{r}"].fill = amarelo
        manual[f"G{r}"].alignment = Alignment(wrap_text=True, vertical="center")
        manual.row_dimensions[r].height = 30
    estados = f"F4:F{m_ultima}"
    for texto, fundo, letra in (("Completa", "verde_claro", "verde"), ("Incompleta", "vermelho_claro", "vermelho"),
                                ("Por preencher", "amarelo", "laranja"), ("Sem medição", "fundo", "suave")):
        manual.conditional_formatting.add(estados, FormulaRule(
            formula=[f'$F4="{texto}"'], fill=PatternFill("solid", bgColor=COR[fundo]), font=Font(color=COR[letra])))
    for col, maximo, nome in (("C", 300, "cintura"), ("E", 150, "pescoço")):
        dv = DataValidation(type="decimal", operator="between", formula1="1", formula2=str(maximo), allow_blank=True,
                            showErrorMessage=True, errorTitle="Valor inválido",
                            error=f"Introduza a {nome} em centímetros (1 a {maximo}).")
        manual.add_data_validation(dv)
        dv.add(f"{col}4:{col}{m_ultima}")
    inicio_semana = 'DATE(VALUE(MID(A4,7,4)),VALUE(MID(A4,4,2)),VALUE(LEFT(A4,2)))'
    dv = DataValidation(type="custom", formula1=f"AND(B4>={inicio_semana},B4<={inicio_semana}+6)",
                        allow_blank=True, showErrorMessage=True, errorTitle="Data fora da semana",
                        error="A data da medição tem de pertencer à semana da coluna A.")
    manual.add_data_validation(dv)
    dv.add(f"B4:B{m_ultima}")
    _tabela(manual, "tblMedidasManuais", f"A3:G{m_ultima}")
    _larguras(manual, [24, 16, 13, 16, 13, 15, 62])
    manual.freeze_panes = "A4"

    # --- Resumos ---
    def resumo(ws, nome_tabela, primeira_col, periodos, extra_cab, extra_formulas, titulo, subtitulo, nota):
        cab = [primeira_col, "Início", "Fim", "Leituras", "Dias com medição", "Período parcial",
               "Peso médio (kg)", "Variação do peso (kg)", "Gordura média (%)", "Variação da gordura (pp)",
               *[c for _, c, _ in METRICAS_RESUMO[2:]], *extra_cab]
        grupos = {c: COR["azul_claro"] for c in cab[6:10]}
        grupos.update({c: COR["verde_claro"] for c in cab[10:10 + len(METRICAS_RESUMO) - 2]})
        grupos.update({c: COR["laranja_claro"] for c in extra_cab})
        _titulo(ws, titulo, len(cab), subtitulo, coluna="B")
        _cabecalho(ws, 3, cab, grupos)
        ult = len(periodos) + 3
        colunas_metricas = ["G", "I"] + [get_column_letter(11 + i) for i in range(len(METRICAS_RESUMO) - 2)]
        for i, p in enumerate(periodos):
            r = i + 4
            ws[f"A{r}"] = p["rotulo"] if primeira_col == "Semana" else p["inicio"]
            ws[f"B{r}"], ws[f"C{r}"] = p["inicio"], p["fim"]
            janela = f'{A},">="&$B{r},{A},"<"&$C{r}+1'
            ws[f"D{r}"] = f"=COUNTIFS({janela})"
            ws[f"E{r}"] = f"=SUMPRODUCT(({Q}>=$B{r})*({Q}<=$C{r})/COUNTIF({Q},{Q}))"
            ws[f"F{r}"] = f'=IF($D{r}=0,"",IF(OR($B{r}<{data_inicio},$C{r}>MAX({Q})),"Sim","Não"))'
            for col, (_, _, campo) in zip(colunas_metricas, METRICAS_RESUMO):
                origem = f"'{FOLHA_CORPO}'!${COLUNA_CAMPO[campo]}$4:${COLUNA_CAMPO[campo]}${ultima}"
                ws[f"{col}{r}"] = f'=IF($D{r}=0,"",AVERAGEIFS({origem},{janela}))'
            for col, base in (("H", "G"), ("J", "I")):
                ws[f"{col}{r}"] = '=""' if i == 0 else f'=IF(OR({base}{r}="",{base}{r - 1}=""),"",{base}{r}-{base}{r - 1})'
            for j, formula in enumerate(extra_formulas):
                ws.cell(r, 10 + len(METRICAS_RESUMO) - 1 + j, formula(r))
        formatos = {"A": "@" if primeira_col == "Semana" else '[$-816]mmmm yyyy', "B": "dd/mm/yyyy",
                    "C": "dd/mm/yyyy", "D": "0", "E": "0", "G": "0.00", "H": '+0.00;-0.00;0.00', "I": "0.0",
                    "J": '+0.0;-0.0;0.0'}
        for i in range(len(METRICAS_RESUMO) - 2):
            formatos[get_column_letter(11 + i)] = "0.00" if i < 5 else "0.0"
        for j, c in enumerate(extra_cab):
            formatos[get_column_letter(10 + len(METRICAS_RESUMO) - 1 + j)] = '+0.0;-0.0;0.0' if "Variação" in c else "0.0"
        _corpo_texto(ws, 4, ult, len(cab))
        for col, formato in formatos.items():
            for r in range(4, ult + 1):
                ws[f"{col}{r}"].number_format = formato
        for r in range(4, ult + 1):
            ws[f"F{r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws.conditional_formatting.add(f"A4:{get_column_letter(len(cab))}{ult}", FormulaRule(
            formula=["$D4=0"], font=Font(color=COR["cinza_ponto"])))
        ws.conditional_formatting.add(f"F4:F{ult}", FormulaRule(
            formula=['$F4="Sim"'], fill=PatternFill("solid", bgColor=COR["amarelo"]), font=Font(color=COR["laranja"])))
        _tabela(ws, nome_tabela, f"A3:{get_column_letter(len(cab))}{ult}")
        _larguras(ws, [24 if primeira_col == "Semana" else 16, 11, 11, 10, 11, 10] + [13.5] * (len(cab) - 6))
        ws.freeze_panes = "B4"
        _nota(ws, f"B{ult + 2}", f"{get_column_letter(min(len(cab), 15))}{ult + 2}", nota)

    def manual_semana(col_manual):
        return lambda r: (f"=IFERROR(IF(INDEX('{FOLHA_MANUAL}'!${col_manual}$4:${col_manual}${m_ultima},"
                          f"MATCH($A{r},'{FOLHA_MANUAL}'!$A$4:$A${m_ultima},0))=\"\",\"\","
                          f"INDEX('{FOLHA_MANUAL}'!${col_manual}$4:${col_manual}${m_ultima},"
                          f"MATCH($A{r},'{FOLHA_MANUAL}'!$A$4:$A${m_ultima},0))),\"\")")

    def manual_mes(col_manual):
        return lambda r: (f"=IFERROR(AVERAGEIFS('{FOLHA_MANUAL}'!${col_manual}$4:${col_manual}${m_ultima},"
                          f"'{FOLHA_MANUAL}'!$B$4:$B${m_ultima},\">=\"&$B{r},"
                          f"'{FOLHA_MANUAL}'!$B$4:$B${m_ultima},\"<\"&$C{r}+1),\"\")")

    nota_periodos = ("Médias aritméticas das leituras aceites do período (não médias diárias). Variação = média do "
                     "período menos a do período anterior; gordura em pontos percentuais. 'Período parcial' = começa "
                     "antes do início da transformação ou termina depois da última leitura.")
    resumo(semanal, "tblResumoSemanal", "Semana", semanais,
           ["Cintura (cm)", "Variação da cintura (cm)", "Pescoço (cm)"],
           [manual_semana("C"), manual_semana("D"), manual_semana("E")],
           "Resumo semanal (domingo a sábado)", atualizado,
           nota_periodos + " Cintura: diferença face à medição anterior disponível.")
    resumo(mensal, "tblResumoMensal", "Mês", mensais, ["Cintura média (cm)", "Pescoço médio (cm)"],
           [manual_mes("C"), manual_mes("E")], "Resumo mensal", atualizado,
           nota_periodos + " Cintura e pescoço: média das medições manuais com data nesse mês.")

    # --- Painel ---
    # Coluna A e coluna Q são margens vazias; o conteúdo vai de B a P.
    _titulo(painel, "", 16, coluna="B")
    painel["B1"] = (f'="Transformação corporal  ·  {inicio_txt} a "&{_data_excel(f"MAX({A})")}')
    hora_ultima = f"INDEX({A},MATCH(MAX({U}),{U},0))"
    painel["B2"] = (f'="Última leitura: "&{_data_excel(hora_ultima)}&" às "&RIGHT("0"&HOUR({hora_ultima}),2)&":"&RIGHT("0"&MINUTE({hora_ultima}),2)'
                    f'&"  ·  "&COUNT({A})&" leituras aceites  ·  {atualizado}"')
    painel["B2"].font = Font(name=FONTE, size=9, italic=True, color=COR["suave"])
    B, D, G = rng("Peso (kg)"), rng("Gordura (%)"), rng("Massa muscular (kg)")
    ultimas7 = lambda v: f'AVERAGEIFS({v},{A},">"&MAX({A})-7)'
    primeiras7 = lambda v: f'AVERAGEIFS({v},{A},"<"&MIN({A})+7)'
    base4s = lambda v: f'AVERAGEIFS({v},{A},">"&MAX({A})-28,{A},"<="&MAX({A})-21)'
    ultima_leitura = lambda v: f"INDEX({v},MATCH(MAX({U}),{U},0))"
    MC, MB, ME = (f"'{FOLHA_MANUAL}'!${c}$4:${c}${m_ultima}" for c in "CBE")

    def manual_ultima(v):
        return f'IF(COUNT({v})=0,"",SUMIFS({v},{MB},_xlfn.MAXIFS({MB},{v},">0")))'

    def manual_variacao(v):
        return (f'IF(COUNT({v})<2,"",SUMIFS({v},{MB},_xlfn.MAXIFS({MB},{v},">0"))'
                f'-SUMIFS({v},{MB},_xlfn.MINIFS({MB},{v},">0")))')

    def manual_data(v, prefixo):
        return f'IF(COUNT({v})=0,"","{prefixo} "&{_data_excel(f"_xlfn.MAXIFS({MB},{v},{chr(34)}>0{chr(34)})")})'

    kg, pp = '0.00 "kg"', '0.0 "%"'
    delta = lambda u: f'+0.00 "{u}";-0.00 "{u}";0.00 "{u}"'
    cartoes = [
        (4, [("Último peso", f"={ultima_leitura(B)}", kg, "azul_claro", '"leitura mais recente"'),
             ("Peso (média 7 dias)", f"={ultimas7(B)}", kg, "azul_claro", '"últimos 7 dias"'),
             ("Peso desde o início", f"={ultimas7(B)}-{primeiras7(B)}", delta("kg"), "fundo", '"média 7 dias: atual − inicial"'),
             ("Ritmo recente", f'=IFERROR(({ultimas7(B)}-{base4s(B)})/3,"")', delta("kg/sem."), "fundo", '"últimas 4 semanas"')]),
        (10, [("Gordura (média 7 dias)", f"={ultimas7(D)}", pp, "laranja_claro", '"estimativa da balança"'),
              ("Gordura desde o início", f"={ultimas7(D)}-{primeiras7(D)}", '+0.0 "pp";-0.0 "pp";0.0 "pp"', "fundo", '"pontos percentuais"'),
              ("Músculo (média 7 dias)", f"={ultimas7(G)}", kg, "verde_claro", '"estimativa da balança"'),
              ("Músculo desde o início", f"={ultimas7(G)}-{primeiras7(G)}", delta("kg"), "fundo", '"média 7 dias: atual − inicial"')]),
        (16, [("Última cintura", f"={manual_ultima(MC)}", '0.0 "cm"', "laranja_claro", manual_data(MC, "medida em")),
              ("Cintura desde o início", f"={manual_variacao(MC)}", '+0.0 "cm";-0.0 "cm";0.0 "cm"', "fundo", '"última − primeira medição"'),
              ("Último pescoço", f"={manual_ultima(ME)}", '0.0 "cm"', "azul_claro", manual_data(ME, "medido em")),
              ("Pescoço desde o início", f"={manual_variacao(ME)}", '+0.0 "cm";-0.0 "cm";0.0 "cm"', "fundo", '"última − primeira medição"')]),
    ]
    for linha, grupo in cartoes:
        for j, (rotulo, formula, formato, cor, nota) in enumerate(grupo):
            c1, c2 = get_column_letter(2 + 4 * j), get_column_letter(4 + 4 * j)
            painel.merge_cells(f"{c1}{linha}:{c2}{linha}")
            painel.merge_cells(f"{c1}{linha + 1}:{c2}{linha + 3}")
            painel.merge_cells(f"{c1}{linha + 4}:{c2}{linha + 4}")
            painel[f"{c1}{linha}"] = rotulo
            painel[f"{c1}{linha + 1}"] = formula
            painel[f"{c1}{linha + 4}"] = f"={nota}"
            fundo = PatternFill("solid", fgColor=COR[cor])
            for r in range(linha, linha + 5):
                for col in range(2 + 4 * j, 5 + 4 * j):
                    c = painel.cell(r, col)
                    c.fill = fundo
                    c.border = Border(left=borda if col == 2 + 4 * j else None, right=borda if col == 4 + 4 * j else None,
                                      top=borda if r == linha else None, bottom=borda if r == linha + 4 else None)
            painel[f"{c1}{linha}"].font = Font(name=FONTE, bold=True, size=11, color=COR["suave"])
            painel[f"{c1}{linha + 1}"].font = Font(name=FONTE, bold=True, size=24, color=COR["texto"])
            painel[f"{c1}{linha + 1}"].number_format = formato
            painel[f"{c1}{linha + 4}"].font = Font(name=FONTE, size=9, italic=True, color=COR["suave"])
            for r in (linha, linha + 1, linha + 4):
                painel[f"{c1}{r}"].alignment = Alignment(horizontal="center", vertical="center")
    # Pensado para um monitor largo (~1900 px) a 100 %: margens A e Q com a largura dos espaços entre cartões.
    _larguras(painel, [ESPACO_LARGURA] + [CARTAO_LARGURA] * 3 + ([ESPACO_LARGURA] + [CARTAO_LARGURA] * 3) * 3
              + [ESPACO_LARGURA])
    for r in range(3, 22):
        painel.row_dimensions[r].height = ESPACO_ALTURA if r in (3, 9, 15, 21) else LINHA_CARTAO
    for _, fim in GRAFICOS_LINHAS:
        painel.row_dimensions[fim + 1].height = ESPACO_ALTURA

    pesos = [l["Weight"] for l in leituras]
    graficos = [
        (grafico_diario(corpo, ultima, "B", "N", "Peso (kg)", COR["azul"], pesos, [m["Weight"] for m in moveis],
                        AMPLITUDE_MINIMA["peso"], "kg"), (1, GRAFICOS_LINHAS[0][0] - 1, 16, GRAFICOS_LINHAS[0][1])),
        (grafico_diario(corpo, ultima, "D", "O", "Gordura corporal (%)", COR["laranja"],
                        [l["BodyFat"] for l in leituras], [m["BodyFat"] for m in moveis], AMPLITUDE_MINIMA["gordura"], "%"), (1, GRAFICOS_LINHAS[1][0] - 1, 16, GRAFICOS_LINHAS[1][1])),
        (grafico_diario(corpo, ultima, "G", "P", "Massa muscular (kg)", COR["verde"],
                        [l["MuscleMass"] for l in leituras], [m["MuscleMass"] for m in moveis], AMPLITUDE_MINIMA["musculo"], "kg"), (1, GRAFICOS_LINHAS[2][0] - 1, 16, GRAFICOS_LINHAS[2][1])),
        (grafico_manual(manual, semanas, "C", "cintura", "Cintura (cm)", COR["laranja"], AMPLITUDE_MINIMA["cintura"]), (1, GRAFICOS_LINHAS[3][0] - 1, 8, GRAFICOS_LINHAS[3][1])),
        (grafico_manual(manual, semanas, "E", "pescoco", "Pescoço (cm)", COR["azul"], AMPLITUDE_MINIMA["pescoco"]), (9, GRAFICOS_LINHAS[3][0] - 1, 16, GRAFICOS_LINHAS[3][1])),
    ]
    for ch, (c0, r0, c1, r1) in graficos:
        if isinstance(ch, ScatterChart):
            _escala_datas(ch.x_axis, _serie(min(l.local for l in leituras)), _serie(max(l.local for l in leituras)))
        # Âncora a duas células: o gráfico ocupa exatamente as colunas/linhas indicadas.
        painel.add_chart(ch, TwoCellAnchor(editAs="oneCell", _from=AnchorMarker(col=c0, row=r0), to=AnchorMarker(col=c1, row=r1)))
    _nota(painel, "B98", "P98", "Gráficos de peso, gordura e músculo: pontos = leituras; linha = média móvel de 7 dias. "
          "Cintura e pescoço: pontos = medições reais; tracejado = semanas sem medição (só visual, sem valores inventados). "
          "As escalas e os tracejados são recalculados em cada atualização.")
    painel.freeze_panes = "A3"

    # --- Critérios e qualidade ---
    _criterios(criterios, cfg, selecao, semanas, atualizado)

    for ws in (mensal, semanal, manual, corpo, criterios):
        ws.sheet_properties.tabColor = None
    painel.sheet_properties.tabColor = COR["azul"]
    manual.sheet_properties.tabColor = "F2C744"
    wb.active = 0
    wb.save(caminho)
    _ignorar_avisos_de_formula(caminho)


def _ignorar_avisos_de_formula(caminho: Path) -> None:
    """O openpyxl não escreve <ignoredErrors>. Sem isto, o Excel marca com triângulos verdes fórmulas
    'inconsistentes' que são intencionais (cada coluna de resumo lê uma métrica diferente)."""
    marca = '<ignoredErrors><ignoredError sqref="A1:Z3000" formula="1" formulaRange="1" numberStoredAsText="1"/></ignoredErrors>'
    temporario = caminho.with_name(caminho.name + ".tmp")
    with zipfile.ZipFile(caminho) as origem, zipfile.ZipFile(temporario, "w", zipfile.ZIP_DEFLATED) as destino:
        for info in origem.infolist():
            dados = origem.read(info.filename)
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", info.filename):
                texto = dados.decode("utf-8")
                # Ordem do esquema: ignoredErrors vem antes de drawing/legacyDrawing/tableParts/extLst.
                m = re.search(r"<(?:drawing|legacyDrawing|tableParts|extLst)\b|</worksheet>", texto)
                texto = texto[:m.start()] + marca + texto[m.start():]
                dados = texto.encode("utf-8")
            destino.writestr(info, dados)
    os.replace(temporario, caminho)


def _criterios(ws, cfg: Config, selecao: Selecao, semanas: list[Semana], atualizado: str) -> None:
    from .leituras import CAMPOS_COMPOSICAO, LIMITES, METRICAS
    leituras = selecao.aceites
    _titulo(ws, "Critérios e qualidade dos dados", 6, atualizado)
    _cabecalho(ws, 3, ["Campo", "Unidade", "Significado", "Cobertura", "Origem", "Campo na exportação"])
    linhas = []
    for m in METRICAS:
        n = sum(1 for l in leituras if isinstance(l[m.campo], (int, float)))
        linhas.append([m.nome, m.unidade, m.significado,
                       f"{n} de {len(leituras)} leituras", m.origem, m.campo])
    linhas += [
        ["Massa gorda", "kg", "Peso × gordura / 100", f"{len(leituras)} leituras", "derivada", "—"],
        ["Médias de 7 dias", "kg / %", "Média das leituras em ]t − 7 dias, t]", f"{len(leituras)} leituras", "derivada", "—"],
        ["Cintura", "cm", "Circunferência ao nível do umbigo (média de 3 leituras)",
         f"{sum(s.cintura is not None for s in semanas)} de {len(semanas)} semanas", "manual", "—"],
        ["Pescoço", "cm", "Circunferência do pescoço (média de 3 leituras)",
         f"{sum(s.pescoco is not None for s in semanas)} de {len(semanas)} semanas", "manual", "—"],
    ]
    for i, linha in enumerate(linhas, start=4):
        for j, v in enumerate(linha, start=1):
            ws.cell(i, j, v)
    fim = 3 + len(linhas)
    _corpo_texto(ws, 4, fim, 6)
    for r in range(4, fim + 1):
        for c in range(1, 7):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="center")
            ws.cell(r, c).border = Border(bottom=borda)
    c = selecao.contagens
    regras = [
        ("Início", f"Leituras com data local a partir de {cfg.inicio:%d/%m/%Y} ({cfg.fuso_nome})."),
        ("Protocolo de pesagem", cfg.protocolo_pesagem or "Não indicado."),
        ("Perfil", f"Titular: identificador Zepp {cfg.perfil_titular_id} (o leitor pede apenas este membro). "
                   + (f"Exportações antigas sem identificador: altura {cfg.altura_perfil_sem_id_cm:g} cm."
                      if cfg.altura_perfil_sem_id_cm else "")),
        ("Composição completa", "IMC, gordura, água, massa óssea, metabolismo, músculo, visceral, proteína, "
                                "massa proteica, idade corporal e pontuação com valores numéricos."),
        ("Plausibilidade", "; ".join(f"{k} {lo}–{hi}" for k, (lo, hi) in LIMITES.items())),
        ("Semanas", "Domingo a sábado, em hora local."),
        ("Médias", "Por leitura (não por dia). 'Dias com medição' mostra quantos dias distintos contribuíram."),
        ("Contagens", f"origem {c['origem']} · anteriores ao início {c['anteriores_ao_inicio']} · "
                      f"incompletas {c['incompletas']} · implausíveis {c['implausiveis']} · "
                      f"outro perfil {c['outro_perfil']} · aceites {c['aceites']}"),
        ("Limites", "Composição, idade e pontuação são estimativas da aplicação, não diagnósticos clínicos."),
        ("Atualização", "Sincronizar a Zepp Life, guardar e fechar o Excel e executar ATUALIZAR-EXCEL.cmd. "
                        "Depois de editar só medidas manuais: RECALCULAR-SEM-DESCARREGAR.cmd."),
    ]
    r = fim + 2
    ws.cell(r, 1, "Regras").font = Font(name=FONTE, bold=True, size=12, color=COR["titulo"])
    for nome, texto in regras:
        r += 1
        ws.cell(r, 1, nome).font = Font(name=FONTE, bold=True, size=10, color=COR["texto"])
        ws.cell(r, 1).alignment = Alignment(vertical="center")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.cell(r, 2, texto).font = Font(name=FONTE, size=10, color=COR["texto"])
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 30
    r += 2
    ws.cell(r, 1, f"Pesagens excluídas desde o início ({len(selecao.excluidas)})").font = \
        Font(name=FONTE, bold=True, size=12, color=COR["titulo"])
    r += 1
    _cabecalho(ws, r, ["Data/hora", "Peso (kg)", "Motivo"])
    ws.row_dimensions[r].height = 24
    for leitura, motivo in selecao.excluidas:
        r += 1
        ws.cell(r, 1, leitura.local).number_format = "dd/mm/yyyy hh:mm"
        ws.cell(r, 2, leitura["Weight"]).number_format = "0.00"
        ws.cell(r, 3, motivo)
        for col in range(1, 4):
            ws.cell(r, col).font = Font(name=FONTE, size=10, color=COR["texto"])
    _larguras(ws, [24, 12, 46, 22, 14, 20])
    ws.freeze_panes = "A4"
