"""Folha "Painel": cartões com indicadores e gráficos.

Desenhado para o ecrã do utilizador (janela com ~1858 px úteis, 1 carácter de largura ≈ 9 px,
1 ponto de altura ≈ 1,667 px). Um único espaçamento (~27 px) em todo o lado: margens laterais,
espaços entre cartões (horizontais e verticais), antes dos gráficos e entre gráficos.
Colunas A e Q são margens vazias; o conteúdo vai de B a P (4 cartões de 3 colunas + 3 espaços).
"""
from __future__ import annotations

from dataclasses import dataclass

from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Alignment, Border, PatternFill
from openpyxl.utils import get_column_letter

from ..resumos import Calculos
from ..verificacao import Esperado
from . import estilos, graficos
from .estilos import BORDA, COR
from .folha_manuais import FolhaManuais
from .folha_medicoes import FolhaMedicoes

NOME = "Painel"
ESPACO_LARGURA = 3          # caracteres ≈ 27 px
ESPACO_ALTURA = 16.2        # pontos ≈ 27 px
CARTAO_LARGURA = 15.9       # 3 colunas por cartão: 4 cartões + 5 espaços ≈ 1851 px
LINHA_CARTAO = 17           # pontos
LINHAS_CARTOES = [4, 10, 16]                               # cada fila: rótulo, valor (3 linhas), nota
GRAFICOS_LINHAS = [(22, 39), (41, 58), (60, 77), (79, 96)]  # 18 linhas cada, um espaço depois
LINHA_NOTA = 98
ULTIMA_COLUNA = 16                                         # P
# Amplitude mínima dos eixos Y (unidade da métrica): evita ampliar ruído de medição.
AMPLITUDE_MINIMA = {"peso": 6, "gordura": 4, "musculo": 3, "cintura": 10, "pescoco": 3}

KG, PERCENTAGEM, CM = '0.00 "kg"', '0.0 "%"', '0.0 "cm"'


def _delta(unidade: str, casas: int = 2) -> str:
    n = "0." + "0" * casas
    return f'+{n} "{unidade}";-{n} "{unidade}";{n} "{unidade}"'


@dataclass
class Cartao:
    rotulo: str
    formula: str
    formato: str
    cor: str
    nota: str               # expressão Excel (texto entre aspas ou fórmula)
    esperado: object


def _cartoes(calc: Calculos, medicoes: FolhaMedicoes, manuais: FolhaManuais) -> list[list[Cartao]]:
    datas, instantes = medicoes.intervalo("Data/hora"), medicoes.intervalo("Instante UTC")
    peso, gordura, musculo = (medicoes.intervalo_metrica(c) for c in ("Weight", "BodyFat", "MuscleMass"))
    ultimas7 = lambda v: f'AVERAGEIFS({v},{datas},">"&MAX({datas})-7)'
    primeiras7 = lambda v: f'AVERAGEIFS({v},{datas},"<"&MIN({datas})+7)'
    base4s = lambda v: f'AVERAGEIFS({v},{datas},">"&MAX({datas})-28,{datas},"<="&MAX({datas})-21)'
    # A última leitura é a de maior instante UTC (não de maior hora local: mudança de hora).
    ultima = lambda v: f"INDEX({v},MATCH(MAX({instantes}),{instantes},0))"
    datas_m = manuais.intervalo("Data da medição")

    def manual_ultima(v):
        return f'IF(COUNT({v})=0,"",SUMIFS({v},{datas_m},_xlfn.MAXIFS({datas_m},{v},">0")))'

    def manual_variacao(v):
        return (f'IF(COUNT({v})<2,"",SUMIFS({v},{datas_m},_xlfn.MAXIFS({datas_m},{v},">0"))'
                f'-SUMIFS({v},{datas_m},_xlfn.MINIFS({datas_m},{v},">0")))')

    def manual_data(v, prefixo):
        data = f'_xlfn.MAXIFS({datas_m},{v},">0")'
        return f'IF(COUNT({v})=0,"","{prefixo} "&{estilos.data_formula(data)})'

    ind = calc.indicadores
    p, g, m = ind["peso_kg"], ind["gordura_corporal_percentagem"], ind["massa_muscular_kg"]
    cintura, pescoco = ind["cintura_cm"], ind["pescoco_cm"]
    variacao_manual = lambda i: i.variacao if i and i.medicoes > 1 else None
    media7 = '"média 7 dias: atual − inicial"'
    v_cintura, v_pescoco = manuais.intervalo("Cintura (cm)"), manuais.intervalo("Pescoço (cm)")
    return [
        [Cartao("Último peso", f"={ultima(peso)}", KG, "azul_claro", '"leitura mais recente"', p.ultima_leitura),
         Cartao("Peso (média 7 dias)", f"={ultimas7(peso)}", KG, "azul_claro", '"últimos 7 dias"', p.atual_media_7d),
         Cartao("Peso desde o início", f"={ultimas7(peso)}-{primeiras7(peso)}", _delta("kg"), "fundo", media7,
                p.variacao),
         Cartao("Ritmo recente", f'=IFERROR(({ultimas7(peso)}-{base4s(peso)})/3,"")', _delta("kg/sem."), "fundo",
                '"últimas 4 semanas"', p.ritmo_semanal_4s)],
        [Cartao("Gordura (média 7 dias)", f"={ultimas7(gordura)}", PERCENTAGEM, "laranja_claro",
                '"estimativa da balança"', g.atual_media_7d),
         Cartao("Gordura desde o início", f"={ultimas7(gordura)}-{primeiras7(gordura)}", _delta("pp", 1), "fundo",
                '"pontos percentuais"', g.variacao),
         Cartao("Músculo (média 7 dias)", f"={ultimas7(musculo)}", KG, "verde_claro", '"estimativa da balança"',
                m.atual_media_7d),
         Cartao("Músculo desde o início", f"={ultimas7(musculo)}-{primeiras7(musculo)}", _delta("kg"), "fundo", media7,
                m.variacao)],
        [Cartao("Última cintura", f"={manual_ultima(v_cintura)}", CM, "laranja_claro",
                manual_data(v_cintura, "medida em"), cintura.ultima if cintura else None),
         Cartao("Cintura desde o início", f"={manual_variacao(v_cintura)}", _delta("cm", 1), "fundo",
                '"última − primeira medição"', variacao_manual(cintura)),
         Cartao("Último pescoço", f"={manual_ultima(v_pescoco)}", CM, "azul_claro",
                manual_data(v_pescoco, "medido em"), pescoco.ultima if pescoco else None),
         Cartao("Pescoço desde o início", f"={manual_variacao(v_pescoco)}", _delta("cm", 1), "fundo",
                '"última − primeira medição"', variacao_manual(pescoco))],
    ]


def _desenhar_cartao(ws, linha: int, indice: int, c: Cartao) -> str:
    """Desenha um cartão (rótulo, valor em 3 linhas, nota) e devolve a célula do valor."""
    inicio, fim = 2 + 4 * indice, 4 + 4 * indice
    c1, c2 = get_column_letter(inicio), get_column_letter(fim)
    ws.merge_cells(f"{c1}{linha}:{c2}{linha}")
    ws.merge_cells(f"{c1}{linha + 1}:{c2}{linha + 3}")
    ws.merge_cells(f"{c1}{linha + 4}:{c2}{linha + 4}")
    ws[f"{c1}{linha}"] = c.rotulo
    ws[f"{c1}{linha + 1}"] = c.formula
    ws[f"{c1}{linha + 4}"] = f"={c.nota}"
    fundo = PatternFill("solid", fgColor=COR[c.cor])
    for r in range(linha, linha + 5):
        for col in range(inicio, fim + 1):
            celula = ws.cell(r, col)
            celula.fill = fundo
            celula.border = Border(left=BORDA if col == inicio else None, right=BORDA if col == fim else None,
                                   top=BORDA if r == linha else None, bottom=BORDA if r == linha + 4 else None)
    ws[f"{c1}{linha}"].font = estilos.fonte(11, "suave", negrito=True)
    ws[f"{c1}{linha + 1}"].font = estilos.fonte(24, negrito=True)
    ws[f"{c1}{linha + 1}"].number_format = c.formato
    ws[f"{c1}{linha + 4}"].font = estilos.fonte(9, "suave", italico=True)
    for r in (linha, linha + 1, linha + 4):
        ws[f"{c1}{r}"].alignment = Alignment(horizontal="center", vertical="center")
    return f"{c1}{linha + 1}"


def _graficos(calc: Calculos, medicoes: FolhaMedicoes, manuais: FolhaManuais) -> list:
    leituras, moveis, semanas = calc.leituras, calc.moveis, calc.semanas
    inicio, fim = min(lt.local for lt in leituras), max(lt.local for lt in leituras)
    x = medicoes.intervalo("Data/hora")

    def diario(campo, titulo, cor, amplitude, unidade, media):
        return graficos.diario(x, medicoes.intervalo_metrica(campo), medicoes.intervalo(media), titulo, cor,
                               [lt[campo] for lt in leituras], [m[campo] for m in moveis], inicio, fim, amplitude, unidade)

    datas = [s.data for s in semanas if s.data is not None] or [s.inicio for s in semanas]

    def manual(cabecalho, campo, titulo, cor, amplitude):
        return graficos.manual(manuais.intervalo(cabecalho), manuais.intervalo("Data da medição"),
                               [getattr(s, campo) for s in semanas], datas, titulo, cor, amplitude)

    (l1, f1), (l2, f2), (l3, f3), (l4, f4) = GRAFICOS_LINHAS
    # (gráfico, (coluna inicial, linha inicial, coluna final, linha final)), índices a começar em 0
    return [
        (diario("Weight", "Peso (kg)", COR["azul"], AMPLITUDE_MINIMA["peso"], "kg", "Peso, média 7 dias (kg)"),
         (1, l1 - 1, ULTIMA_COLUNA, f1)),
        (diario("BodyFat", "Gordura corporal (%)", COR["laranja"], AMPLITUDE_MINIMA["gordura"], "%",
                "Gordura, média 7 dias (%)"), (1, l2 - 1, ULTIMA_COLUNA, f2)),
        (diario("MuscleMass", "Massa muscular (kg)", COR["verde"], AMPLITUDE_MINIMA["musculo"], "kg",
                "Músculo, média 7 dias (kg)"), (1, l3 - 1, ULTIMA_COLUNA, f3)),
        (manual("Cintura (cm)", "cintura", "Cintura (cm)", COR["laranja"], AMPLITUDE_MINIMA["cintura"]),
         (1, l4 - 1, 8, f4)),
        (manual("Pescoço (cm)", "pescoco", "Pescoço (cm)", COR["azul"], AMPLITUDE_MINIMA["pescoco"]),
         (9, l4 - 1, ULTIMA_COLUNA, f4)),
    ]


def construir(ws, calc: Calculos, medicoes: FolhaMedicoes, manuais: FolhaManuais, inicio_txt: str,
              atualizado: str, esperados: list[Esperado]) -> None:
    datas, instantes = medicoes.intervalo("Data/hora"), medicoes.intervalo("Instante UTC")
    estilos.titulo(ws, "", ULTIMA_COLUNA, coluna="B")
    ws["B1"] = f'="Transformação corporal  ·  {inicio_txt} a "&{estilos.data_formula(f"MAX({datas})")}'
    hora = f"INDEX({datas},MATCH(MAX({instantes}),{instantes},0))"
    ws["B2"] = (f'="Última leitura: "&{estilos.data_formula(hora)}&" às "&RIGHT("0"&HOUR({hora}),2)&":"'
                f'&RIGHT("0"&MINUTE({hora}),2)&"  ·  "&COUNT({datas})&" leituras aceites  ·  {atualizado}"')
    ws["B2"].font = estilos.fonte(9, "suave", italico=True)
    for linha, fila in zip(LINHAS_CARTOES, _cartoes(calc, medicoes, manuais), strict=True):
        for indice, cartao in enumerate(fila):
            celula = _desenhar_cartao(ws, linha, indice, cartao)
            esperados.append(Esperado(NOME, celula, cartao.esperado, cartao.rotulo))
    estilos.larguras(ws, [ESPACO_LARGURA] + [CARTAO_LARGURA] * 3 + ([ESPACO_LARGURA] + [CARTAO_LARGURA] * 3) * 3
                     + [ESPACO_LARGURA])
    espacos = {linha - 1 for linha in LINHAS_CARTOES} | {LINHAS_CARTOES[-1] + 5}
    for r in range(LINHAS_CARTOES[0] - 1, LINHAS_CARTOES[-1] + 6):
        ws.row_dimensions[r].height = ESPACO_ALTURA if r in espacos else LINHA_CARTAO
    for _, fim in GRAFICOS_LINHAS:
        ws.row_dimensions[fim + 1].height = ESPACO_ALTURA
    for grafico, (c0, r0, c1, r1) in _graficos(calc, medicoes, manuais):
        # Âncora a duas células: o gráfico ocupa exatamente as colunas/linhas indicadas.
        ws.add_chart(grafico, TwoCellAnchor(editAs="oneCell", _from=AnchorMarker(col=c0, row=r0),
                                            to=AnchorMarker(col=c1, row=r1)))
    estilos.nota(ws, f"B{LINHA_NOTA}", f"{get_column_letter(ULTIMA_COLUNA)}{LINHA_NOTA}",
                 "Gráficos de peso, gordura e músculo: pontos = leituras; linha = média móvel de 7 dias. "
                 "Cintura e pescoço: pontos = medições reais; tracejado = semanas sem medição (só visual, sem "
                 "valores inventados). As escalas e os tracejados são recalculados em cada atualização.")
    ws.freeze_panes = "A3"
