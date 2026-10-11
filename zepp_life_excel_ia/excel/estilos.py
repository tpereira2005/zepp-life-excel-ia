"""Cores, tipos de letra e blocos de formatação comuns às folhas."""
from __future__ import annotations

import os
import re
import zipfile
from pathlib import Path

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONTE = "Arial"
COR = {
    "titulo": "243447", "texto": "263238", "suave": "52616B", "linha": "D7DEE3", "fundo": "F4F7F9",
    "cabecalho": "E8EDF3", "azul": "2F6BFF", "azul_claro": "EAF0FF", "laranja": "C76D00",
    "laranja_claro": "FFF3D6", "verde": "2E7D5B", "verde_claro": "E5F5EC", "vermelho": "B9382E",
    "vermelho_claro": "FCE8E6", "amarelo": "FFF7D6", "cinza_ponto": "9AA7B0",
}
BORDA = Side(style="thin", color=COR["linha"])


def fonte(tamanho: int = 10, cor: str = "texto", negrito: bool = False, italico: bool = False) -> Font:
    # None (e não False) quando não se pede: o XML fica sem os atributos <b>/<i>.
    return Font(name=FONTE, size=tamanho, bold=negrito or None, italic=italico or None, color=COR[cor])


def titulo(ws, texto: str, ultima_coluna: int, subtitulo: str | None = None, coluna: str = "A") -> None:
    """coluna = onde começam título e subtítulo ("B" quando a coluna A está fixa, para a linha
    de fixação do Excel não atravessar o texto)."""
    ws.sheet_view.showGridLines = False
    fim = get_column_letter(ultima_coluna)
    ws.merge_cells(f"{coluna}1:{fim}1")
    ws[f"{coluna}1"] = texto
    ws[f"{coluna}1"].font = fonte(16, "titulo", negrito=True)
    ws[f"{coluna}1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    if subtitulo:
        ws.merge_cells(f"{coluna}2:{fim}2")
        ws[f"{coluna}2"] = subtitulo
        ws[f"{coluna}2"].font = fonte(9, "suave", italico=True)
        ws[f"{coluna}2"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 20


def cabecalho(ws, linha: int, nomes: list[str], grupos: dict[str, str] | None = None) -> None:
    """grupos: cabeçalho -> cor de fundo (hex) para destacar grupos de colunas."""
    for i, nome in enumerate(nomes, start=1):
        c = ws.cell(linha, i, nome)
        c.font = fonte(negrito=True)
        c.fill = PatternFill("solid", fgColor=(grupos or {}).get(nome, COR["cabecalho"]))
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=BORDA)
    ws.row_dimensions[linha].height = 58


def tabela(ws, nome: str, ref: str) -> None:
    t = Table(displayName=nome, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)


def corpo_texto(ws, inicio_linha: int, fim_linha: int, colunas: int) -> None:
    """Tipo de letra e alinhamento vertical das células de dados (mantém o alinhamento horizontal)."""
    for linha in ws.iter_rows(min_row=inicio_linha, max_row=fim_linha, max_col=colunas):
        for c in linha:
            c.font = fonte()
            c.alignment = Alignment(vertical="center", horizontal=c.alignment.horizontal,
                                    wrap_text=c.alignment.wrap_text)


def larguras(ws, valores: list[float]) -> None:
    for i, largura in enumerate(valores, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura


def nota(ws, celula: str, fim: str, texto: str) -> None:
    ws.merge_cells(f"{celula}:{fim}")
    ws[celula] = texto
    ws[celula].font = fonte(9, "suave", italico=True)
    ws[celula].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[ws[celula].row].height = 30


def data_formula(expr: str) -> str:
    """dd/mm/aaaa numa fórmula. TEXT() depende do idioma do Excel (aaaa vs yyyy); DAY/MONTH/YEAR não."""
    return f'RIGHT("0"&DAY({expr}),2)&"/"&RIGHT("0"&MONTH({expr}),2)&"/"&YEAR({expr})'


def intervalo(folha: str, coluna: str, primeira: int, ultima: int) -> str:
    """Referência absoluta a uma coluna de outra folha, ex.: 'Medidas Manuais'!$C$4:$C$40."""
    return f"'{folha}'!${coluna}${primeira}:${coluna}${ultima}"


def ignorar_avisos_de_formula(caminho: Path) -> None:
    """O openpyxl não escreve <ignoredErrors>. Sem isto, o Excel marca com triângulos verdes fórmulas
    'inconsistentes' que são intencionais (cada coluna de resumo lê uma métrica diferente)."""
    marca = ('<ignoredErrors><ignoredError sqref="A1:Z3000" formula="1" formulaRange="1" '
             'numberStoredAsText="1"/></ignoredErrors>')
    temporario = caminho.with_name(caminho.name + ".tmp")
    with zipfile.ZipFile(caminho) as origem, zipfile.ZipFile(temporario, "w", zipfile.ZIP_DEFLATED) as destino:
        for info in origem.infolist():
            dados = origem.read(info.filename)
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", info.filename):
                texto = dados.decode("utf-8")
                # Ordem do esquema: ignoredErrors vem antes de drawing/legacyDrawing/tableParts/extLst.
                m = re.search(r"<(?:drawing|legacyDrawing|tableParts|extLst)\b|</worksheet>", texto)
                texto = texto[:m.start()] + marca + texto[m.start():]
                dados = texto.encode("utf-8")
            destino.writestr(info, dados)
    os.replace(temporario, caminho)
