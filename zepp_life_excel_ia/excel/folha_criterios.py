"""Folha "Critérios e Qualidade": dicionário de campos, regras de seleção e pesagens excluídas."""
from __future__ import annotations

from openpyxl.styles import Alignment, Border

from ..metricas import DA_EXPORTACAO, LIMITES
from ..resumos import Calculos
from . import estilos
from .estilos import BORDA


def _regras(calc: Calculos) -> list[tuple[str, str]]:
    cfg, c = calc.cfg, calc.selecao.contagens
    perfil = f"Titular: identificador Zepp {cfg.perfil_titular_id} (o leitor pede apenas este membro). "
    if cfg.altura_perfil_sem_id_cm:
        perfil += f"Exportações antigas sem identificador: altura {cfg.altura_perfil_sem_id_cm:g} cm."
    return [
        ("Início", f"Leituras com data local a partir de {cfg.inicio:%d/%m/%Y} ({cfg.fuso_nome})."),
        ("Protocolo de pesagem", cfg.protocolo_pesagem or "Não indicado."),
        ("Perfil", perfil),
        ("Composição completa", "IMC, gordura, água, massa óssea, metabolismo, músculo, visceral, proteína, "
                                "massa proteica, idade corporal e pontuação com valores numéricos."),
        ("Plausibilidade", "; ".join(f"{k} {lo}–{hi}" for k, (lo, hi) in LIMITES.items())),
        ("Semanas", "Domingo a sábado, em hora local."),
        ("Médias", "Por leitura (não por dia). 'Dias com medição' mostra quantos dias distintos contribuíram."),
        ("Contagens", f"origem {c['origem']} · anteriores ao início {c['anteriores_ao_inicio']} · "
                      f"incompletas {c['incompletas']} · implausíveis {c['implausiveis']} · "
                      f"outro perfil {c['outro_perfil']} · aceites {c['aceites']}"),
        ("Limites", "Composição, idade e pontuação são estimativas da aplicação, não diagnósticos clínicos."),
        ("Atualização", "Sincronizar a Zepp Life, guardar e fechar o Excel e executar ATUALIZAR-EXCEL.cmd. "
                        "Depois de editar só medidas manuais: RECALCULAR-SEM-DESCARREGAR.cmd."),
    ]


def construir(ws, calc: Calculos, atualizado: str) -> None:
    leituras, semanas, selecao = calc.leituras, calc.semanas, calc.selecao
    estilos.titulo(ws, "Critérios e qualidade dos dados", 6, atualizado)
    estilos.cabecalho(ws, 3, ["Campo", "Unidade", "Significado", "Cobertura", "Origem", "Campo na exportação"])
    linhas = [[m.nome, m.unidade, m.significado,
               f"{sum(isinstance(lt[m.campo], (int, float)) for lt in leituras)} de {len(leituras)} leituras",
               m.origem, m.campo] for m in DA_EXPORTACAO]
    linhas += [
        ["Massa gorda", "kg", "Peso × gordura / 100", f"{len(leituras)} leituras", "derivada", "—"],
        ["Médias de 7 dias", "kg / %", "Média das leituras em ]t − 7 dias, t]", f"{len(leituras)} leituras",
         "derivada", "—"],
        ["Cintura", "cm", "Circunferência ao nível do umbigo (média de 3 leituras)",
         f"{sum(s.cintura is not None for s in semanas)} de {len(semanas)} semanas", "manual", "—"],
        ["Pescoço", "cm", "Circunferência do pescoço (média de 3 leituras)",
         f"{sum(s.pescoco is not None for s in semanas)} de {len(semanas)} semanas", "manual", "—"],
    ]
    for i, linha in enumerate(linhas, start=4):
        for j, v in enumerate(linha, start=1):
            ws.cell(i, j, v)
    fim = 3 + len(linhas)
    estilos.corpo_texto(ws, 4, fim, 6)
    for r in range(4, fim + 1):
        for c in range(1, 7):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="center")
            ws.cell(r, c).border = Border(bottom=BORDA)
    r = fim + 2
    ws.cell(r, 1, "Regras").font = estilos.fonte(12, "titulo", negrito=True)
    for nome, texto in _regras(calc):
        r += 1
        ws.cell(r, 1, nome).font = estilos.fonte(negrito=True)
        ws.cell(r, 1).alignment = Alignment(vertical="center")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.cell(r, 2, texto).font = estilos.fonte()
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 30
    r += 2
    ws.cell(r, 1, f"Pesagens excluídas desde o início ({len(selecao.excluidas)})").font = \
        estilos.fonte(12, "titulo", negrito=True)
    r += 1
    estilos.cabecalho(ws, r, ["Data/hora", "Peso (kg)", "Motivo"])
    ws.row_dimensions[r].height = 24
    for leitura, motivo in selecao.excluidas:
        r += 1
        ws.cell(r, 1, leitura.local).number_format = "dd/mm/yyyy hh:mm"
        ws.cell(r, 2, leitura["Weight"]).number_format = "0.00"
        ws.cell(r, 3, motivo)
        for col in range(1, 4):
            ws.cell(r, col).font = estilos.fonte()
    estilos.larguras(ws, [24, 12, 46, 22, 14, 20])
    ws.freeze_panes = "A4"
