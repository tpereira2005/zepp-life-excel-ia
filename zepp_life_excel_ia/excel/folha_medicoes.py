"""Folha "Medições Corporais": uma linha por leitura aceite, com fórmulas de massa gorda e médias de 7 dias."""
from __future__ import annotations

from dataclasses import dataclass

from openpyxl.utils import get_column_letter

from ..metricas import COLUNAS_MEDICOES, MEDIAS_MOVEIS, POR_CAMPO
from ..resumos import Calculos
from ..verificacao import Esperado
from . import estilos

NOME = "Medições Corporais"
PRIMEIRA = 4   # primeira linha de dados

# (cabeçalho, largura, formato). "Dia" e "Instante UTC" ficam ocultas: servem as fórmulas.
COLUNAS = [
    ("Data/hora", 19, "dd/mm/yyyy hh:mm"),
    *[(POR_CAMPO[c].coluna, POR_CAMPO[c].largura, POR_CAMPO[c].formato) for c in COLUNAS_MEDICOES],
    *[(m.coluna, 14, "0.00") for m in MEDIAS_MOVEIS],
    ("Dia", 11, "dd/mm/yyyy"),
    # Instante UTC: identifica a última pesagem mesmo quando duas têm a mesma hora local
    # (mudança de hora de outubro) e mesmo que a tabela seja reordenada no Excel.
    ("Instante UTC", 19, "dd/mm/yyyy hh:mm:ss"),
]
OCULTAS = ("Dia", "Instante UTC")
LETRA = {nome: get_column_letter(i) for i, (nome, _, _) in enumerate(COLUNAS, start=1)}


@dataclass
class FolhaMedicoes:
    """Dá às outras folhas referências às colunas desta, sem letras escritas à mão."""
    ultima: int

    def intervalo(self, cabecalho: str) -> str:
        return estilos.intervalo(NOME, LETRA[cabecalho], PRIMEIRA, self.ultima)

    def intervalo_metrica(self, campo: str) -> str:
        return self.intervalo(POR_CAMPO[campo].coluna)


def construir(ws, calc: Calculos, inicio_txt: str, atualizado: str, esperados: list[Esperado]) -> FolhaMedicoes:
    leituras = calc.leituras
    ultima = len(leituras) + PRIMEIRA - 1
    estilos.titulo(ws, f"Medições corporais completas desde {inicio_txt}", len(COLUNAS),
                   f"{len(leituras)} leituras aceites · {atualizado} · fórmulas: massa gorda e médias móveis de 7 dias")
    estilos.cabecalho(ws, PRIMEIRA - 1, [nome for nome, _, _ in COLUNAS])
    data, peso, gordura = LETRA["Data/hora"], LETRA[POR_CAMPO["Weight"].coluna], LETRA[POR_CAMPO["BodyFat"].coluna]
    gorda = LETRA[POR_CAMPO["massa_gorda"].coluna]
    datas = f"${data}${PRIMEIRA}:${data}${ultima}"
    for r, (lt, moveis) in enumerate(zip(leituras, calc.moveis, strict=True), start=PRIMEIRA):
        ws[f"{data}{r}"] = lt.local
        for campo in COLUNAS_MEDICOES:
            if POR_CAMPO[campo].da_exportacao:
                ws[f"{LETRA[POR_CAMPO[campo].coluna]}{r}"] = lt[campo]
        ws[f"{gorda}{r}"] = f'=IF({gordura}{r}="","",{peso}{r}*{gordura}{r}/100)'
        esperados.append(Esperado(NOME, f"{gorda}{r}", lt.massa_gorda, "massa gorda"))
        for m in MEDIAS_MOVEIS:
            origem, destino = LETRA[POR_CAMPO[m.campo].coluna], LETRA[m.coluna]
            ws[f"{destino}{r}"] = (f'=AVERAGEIFS(${origem}${PRIMEIRA}:${origem}${ultima},{datas},'
                                   f'">"&${data}{r}-7,{datas},"<="&${data}{r})')
            esperados.append(Esperado(NOME, f"{destino}{r}", moveis[m.campo], m.coluna))
        ws[f"{LETRA['Dia']}{r}"] = f"=INT({data}{r})"
        ws[f"{LETRA['Instante UTC']}{r}"] = lt.instante_utc.replace(tzinfo=None)
    for i, (_, largura, formato) in enumerate(COLUNAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
        for r in range(PRIMEIRA, ultima + 1):
            ws.cell(r, i).number_format = formato
    estilos.corpo_texto(ws, PRIMEIRA, ultima, len(COLUNAS))
    estilos.tabela(ws, "tblCorpo", f"A{PRIMEIRA - 1}:{get_column_letter(len(COLUNAS))}{ultima}")
    for nome in OCULTAS:
        ws.column_dimensions[LETRA[nome]].hidden = True
    ws.freeze_panes = f"A{PRIMEIRA}"
    return FolhaMedicoes(ultima)
