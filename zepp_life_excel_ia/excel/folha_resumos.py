"""Folhas "Resumo Semanal" e "Resumo Mensal".

Cada coluna é definida num único sítio (`Coluna`): cabeçalho, fórmula, formato e o valor que o Python
espera — este último é usado pela verificação no Excel real.
"""
from __future__ import annotations

from dataclasses import dataclass
from typing import Callable

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..config import Config
from ..metricas import POR_CAMPO, RESUMO
from ..resumos import Periodo
from ..verificacao import Esperado
from . import estilos
from .estilos import COR
from .folha_manuais import FolhaManuais
from .folha_medicoes import FolhaMedicoes

PRIMEIRA = 4
NOTA = ("Médias aritméticas das leituras aceites do período (não médias diárias). Variação = média do "
        "período menos a do período anterior; gordura em pontos percentuais. 'Período parcial' = começa "
        "antes do início da transformação ou termina depois da última leitura.")
VARIACAO_1 = '+0.0;-0.0;0.0'
VARIACAO_2 = '+0.00;-0.00;0.00'


@dataclass
class Coluna:
    cabecalho: str
    valor: Callable[[int, int, Periodo], object]       # (linha, índice do período, período) -> valor ou fórmula
    formato: str | None = None
    esperado: Callable[[Periodo], object] | None = None
    grupo: str | None = None                            # cor de fundo do cabeçalho
    largura: float = 13.5
    centrado: bool = False


def _colunas(semanal: bool, cfg: Config, medicoes: FolhaMedicoes, manuais: FolhaManuais) -> list[Coluna]:
    datas, dias = medicoes.intervalo("Data/hora"), medicoes.intervalo("Dia")
    data_inicio = f"DATE({cfg.inicio.year},{cfg.inicio.month},{cfg.inicio.day})"
    janela = lambda r: f'{datas},">="&$B{r},{datas},"<"&$C{r}+1'      # leituras entre Início e Fim

    def media(campo):
        origem = medicoes.intervalo_metrica(campo)
        return lambda r, i, p: f'=IF($D{r}=0,"",AVERAGEIFS({origem},{janela(r)}))'

    def variacao(base):     # base = letra da coluna da média
        return lambda r, i, p: '=""' if i == 0 else f'=IF(OR({base}{r}="",{base}{r - 1}=""),"",{base}{r}-{base}{r - 1})'

    def metrica(campo, grupo):
        m = POR_CAMPO[campo]
        return Coluna(m.cabecalho_resumo, media(campo), m.formato_resumo, lambda p: p.medias[campo], grupo)

    colunas = [
        Coluna("Semana", lambda r, i, p: p.rotulo, "@", largura=24) if semanal else
        Coluna("Mês", lambda r, i, p: p.inicio, '[$-816]mmmm yyyy', largura=16),
        Coluna("Início", lambda r, i, p: p.inicio, "dd/mm/yyyy", largura=11),
        Coluna("Fim", lambda r, i, p: p.fim, "dd/mm/yyyy", largura=11),
        Coluna("Leituras", lambda r, i, p: f"=COUNTIFS({janela(r)})", "0", lambda p: p.leituras, largura=10),
        Coluna("Dias com medição", lambda r, i, p: f"=SUMPRODUCT(({dias}>=$B{r})*({dias}<=$C{r})/COUNTIF({dias},{dias}))",
               "0", lambda p: p.dias_com_medicao, largura=11),
        Coluna("Período parcial",
               lambda r, i, p: f'=IF($D{r}=0,"",IF(OR($B{r}<{data_inicio},$C{r}>MAX({dias})),"Sim","Não"))',
               None, lambda p: p.parcial, largura=10, centrado=True),
        metrica("Weight", COR["azul_claro"]),
        Coluna("Variação do peso (kg)", variacao("G"), VARIACAO_2, lambda p: p.variacao_peso_kg, COR["azul_claro"]),
        metrica("BodyFat", COR["azul_claro"]),
        Coluna("Variação da gordura (pp)", variacao("I"), VARIACAO_1, lambda p: p.variacao_gordura_pp, COR["azul_claro"]),
        *[metrica(campo, COR["verde_claro"]) for campo in RESUMO[2:]],
    ]
    laranja = COR["laranja_claro"]
    if semanal:
        def da_semana(cabecalho):
            valores, semanas = manuais.intervalo(cabecalho), manuais.intervalo("Semana")
            procura = f"INDEX({valores},MATCH($A{{r}},{semanas},0))"
            return lambda r, i, p: f'=IFERROR(IF({procura.format(r=r)}="","",{procura.format(r=r)}),"")'
        for cabecalho, chave in (("Cintura (cm)", "cintura_cm"), ("Variação da cintura (cm)", "variacao_cintura_cm"),
                                 ("Pescoço (cm)", "pescoco_cm")):
            formato = VARIACAO_1 if "Variação" in cabecalho else "0.0"
            colunas.append(Coluna(cabecalho, da_semana(cabecalho), formato, lambda p, k=chave: p.manuais[k], laranja))
    else:
        def do_mes(cabecalho):
            valores, datas_m = manuais.intervalo(cabecalho), manuais.intervalo("Data da medição")
            return lambda r, i, p: (f'=IFERROR(AVERAGEIFS({valores},{datas_m},">="&$B{r},'
                                    f'{datas_m},"<"&$C{r}+1),"")')
        for cabecalho, origem, chave in (("Cintura média (cm)", "Cintura (cm)", "cintura_media_cm"),
                                         ("Pescoço médio (cm)", "Pescoço (cm)", "pescoco_medio_cm")):
            colunas.append(Coluna(cabecalho, do_mes(origem), "0.0", lambda p, k=chave: p.manuais[k], laranja))
    # As variações referem as colunas das médias pela letra: confirmar que a ordem não mudou.
    assert [c.cabecalho for c in colunas[6:10:2]] == [POR_CAMPO["Weight"].cabecalho_resumo,
                                                      POR_CAMPO["BodyFat"].cabecalho_resumo]
    return colunas


def construir(ws, periodos: list[Periodo], semanal: bool, cfg: Config, medicoes: FolhaMedicoes,
              manuais: FolhaManuais, atualizado: str, esperados: list[Esperado]) -> None:
    nome = "Resumo Semanal" if semanal else "Resumo Mensal"
    colunas = _colunas(semanal, cfg, medicoes, manuais)
    ultima = len(periodos) + PRIMEIRA - 1
    fim = get_column_letter(len(colunas))
    estilos.titulo(ws, "Resumo semanal (domingo a sábado)" if semanal else "Resumo mensal", len(colunas),
                   atualizado, coluna="B")
    estilos.cabecalho(ws, PRIMEIRA - 1, [c.cabecalho for c in colunas],
                      {c.cabecalho: c.grupo for c in colunas if c.grupo})
    for i, p in enumerate(periodos):
        r = i + PRIMEIRA
        for j, c in enumerate(colunas, start=1):
            ws.cell(r, j, c.valor(r, i, p))
            if c.esperado:
                esperados.append(Esperado(nome, f"{get_column_letter(j)}{r}", c.esperado(p), c.cabecalho))
    estilos.corpo_texto(ws, PRIMEIRA, ultima, len(colunas))
    for j, c in enumerate(colunas, start=1):
        for r in range(PRIMEIRA, ultima + 1):
            if c.formato:
                ws.cell(r, j).number_format = c.formato
            if c.centrado:
                ws.cell(r, j).alignment = Alignment(horizontal="center", vertical="center")
    ws.conditional_formatting.add(f"A{PRIMEIRA}:{fim}{ultima}", FormulaRule(
        formula=[f"$D{PRIMEIRA}=0"], font=Font(color=COR["cinza_ponto"])))
    ws.conditional_formatting.add(f"F{PRIMEIRA}:F{ultima}", FormulaRule(
        formula=[f'$F{PRIMEIRA}="Sim"'], fill=PatternFill("solid", bgColor=COR["amarelo"]),
        font=Font(color=COR["laranja"])))
    estilos.tabela(ws, "tblResumoSemanal" if semanal else "tblResumoMensal", f"A{PRIMEIRA - 1}:{fim}{ultima}")
    estilos.larguras(ws, [c.largura for c in colunas])
    ws.freeze_panes = f"B{PRIMEIRA}"
    texto = NOTA + (" Cintura: diferença face à medição anterior disponível." if semanal else
                    " Cintura e pescoço: média das medições manuais com data nesse mês.")
    estilos.nota(ws, f"B{ultima + 2}", f"{get_column_letter(min(len(colunas), 15))}{ultima + 2}", texto)
