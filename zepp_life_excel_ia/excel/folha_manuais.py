"""Folha "Medidas Manuais": registo semanal editável de cintura e pescoço.

Os nomes da folha e dos cabeçalhos são lidos de volta na atualização seguinte (medidas.ler_existentes):
não os mude sem atualizar `medidas.CABECALHOS`.
"""
from __future__ import annotations

from dataclasses import dataclass

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..medidas import CABECALHOS, ESTADOS_EXCEL
from ..medidas import FOLHA as NOME
from ..resumos import Calculos
from ..verificacao import Esperado
from . import estilos
from .estilos import COR

PRIMEIRA = 4
LETRA = {nome: get_column_letter(i) for i, nome in enumerate(CABECALHOS, start=1)}
SEMANA, DATA, CINTURA = LETRA["Semana"], LETRA["Data da medição"], LETRA["Cintura (cm)"]
VARIACAO, PESCOCO, ESTADO, NOTAS = (LETRA["Variação da cintura (cm)"], LETRA["Pescoço (cm)"], LETRA["Estado"],
                                    LETRA["Observações"])
# Estado -> (fundo, letra) da formatação condicional.
CORES_ESTADO = {"Completa": ("verde_claro", "verde"), "Incompleta": ("vermelho_claro", "vermelho"),
                "Por preencher": ("amarelo", "laranja"), "Sem medição": ("fundo", "suave")}


@dataclass
class FolhaManuais:
    ultima: int

    def intervalo(self, cabecalho: str) -> str:
        return estilos.intervalo(NOME, LETRA[cabecalho], PRIMEIRA, self.ultima)


def construir(ws, calc: Calculos, esperados: list[Esperado]) -> FolhaManuais:
    semanas = calc.semanas
    ultima = len(semanas) + PRIMEIRA - 1
    estilos.titulo(ws, "Medidas manuais semanais", len(CABECALHOS),
                   "Preencha as células amarelas (cintura ao nível do umbigo, pescoço). A data tem de pertencer à "
                   "semana. Guarde e feche o Excel antes de atualizar.")
    estilos.cabecalho(ws, PRIMEIRA - 1, CABECALHOS)
    datas, cinturas = f"{DATA}${PRIMEIRA}:{DATA}${ultima}", f"{CINTURA}${PRIMEIRA}:{CINTURA}${ultima}"
    for r, s in enumerate(semanas, start=PRIMEIRA):
        ws[f"{SEMANA}{r}"] = s.rotulo
        ws[f"{DATA}{r}"] = s.data
        ws[f"{CINTURA}{r}"] = s.cintura
        ws[f"{PESCOCO}{r}"] = s.pescoco
        ws[f"{NOTAS}{r}"] = s.notas
        # Variação = cintura menos a medição anterior disponível (pela data de medição).
        ws[f"{VARIACAO}{r}"] = (
            f'=IF(OR({CINTURA}{r}="",COUNTIFS({datas},"<"&{DATA}{r},{cinturas},">0")=0),"",'
            f'{CINTURA}{r}-SUMIFS({cinturas},{datas},_xlfn.MAXIFS({datas},{datas},"<"&{DATA}{r},{cinturas},">0")))')
        ws[f"{ESTADO}{r}"] = (f'=IF(AND({CINTURA}{r}="",{PESCOCO}{r}=""),IF({NOTAS}{r}<>"","Sem medição","Por preencher"),'
                              f'IF(AND({CINTURA}{r}<>"",{PESCOCO}{r}<>""),"Completa","Incompleta"))')
        esperados.append(Esperado(NOME, f"{VARIACAO}{r}", calc.variacoes_cintura[s.inicio], "variação da cintura"))
        esperados.append(Esperado(NOME, f"{ESTADO}{r}", ESTADOS_EXCEL[s.estado], "estado"))
    estilos.corpo_texto(ws, PRIMEIRA, ultima, len(CABECALHOS))
    amarelo = PatternFill("solid", fgColor=COR["amarelo"])
    for r in range(PRIMEIRA, ultima + 1):
        ws[f"{DATA}{r}"].number_format = "dd/mm/yyyy"
        ws[f"{VARIACAO}{r}"].number_format = '+0.0;-0.0;0.0'
        for col in (CINTURA, PESCOCO):
            ws[f"{col}{r}"].number_format = "0.0"
        for col in (DATA, CINTURA, PESCOCO, NOTAS):     # células editáveis
            ws[f"{col}{r}"].fill = amarelo
        ws[f"{NOTAS}{r}"].alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 30
    estados = f"{ESTADO}{PRIMEIRA}:{ESTADO}{ultima}"
    for texto, (fundo, letra) in CORES_ESTADO.items():
        ws.conditional_formatting.add(estados, FormulaRule(
            formula=[f'${ESTADO}{PRIMEIRA}="{texto}"'], fill=PatternFill("solid", bgColor=COR[fundo]),
            font=Font(color=COR[letra])))
    _validacoes(ws, ultima)
    estilos.tabela(ws, "tblMedidasManuais", f"{SEMANA}{PRIMEIRA - 1}:{NOTAS}{ultima}")
    estilos.larguras(ws, [24, 16, 13, 16, 13, 15, 62])
    ws.freeze_panes = f"A{PRIMEIRA}"
    return FolhaManuais(ultima)


def _validacoes(ws, ultima: int) -> None:
    for col, maximo, nome in ((CINTURA, 300, "cintura"), (PESCOCO, 150, "pescoço")):
        dv = DataValidation(type="decimal", operator="between", formula1="1", formula2=str(maximo), allow_blank=True,
                            showErrorMessage=True, errorTitle="Valor inválido",
                            error=f"Introduza a {nome} em centímetros (1 a {maximo}).")
        ws.add_data_validation(dv)
        dv.add(f"{col}{PRIMEIRA}:{col}{ultima}")
    a, b = f"{SEMANA}{PRIMEIRA}", f"{DATA}{PRIMEIRA}"
    inicio_semana = f'DATE(VALUE(MID({a},7,4)),VALUE(MID({a},4,2)),VALUE(LEFT({a},2)))'
    dv = DataValidation(type="custom", formula1=f"AND({b}>={inicio_semana},{b}<={inicio_semana}+6)",
                        allow_blank=True, showErrorMessage=True, errorTitle="Data fora da semana",
                        error="A data da medição tem de pertencer à semana da coluna A.")
    ws.add_data_validation(dv)
    dv.add(f"{DATA}{PRIMEIRA}:{DATA}{ultima}")
