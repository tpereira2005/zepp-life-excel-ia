r"""Compara dois pares (Excel, JSON) gerados — para confirmar que uma alteração ao código não muda o resultado.

Compara cada célula (valor, fórmula, formato, letra, preenchimento, bordas, alinhamento), células unidas,
larguras, alturas, painéis fixos, tabelas, validações, formatação condicional, o XML dos gráficos e dos
desenhos, e o JSON completo. Ignora só o que depende do momento ou do nome: "Atualizado em …",
gerado_em_utc, sha256_excel, gerador e ficheiro_excel.

Uso: .venv\Scripts\python scripts\comparar_saidas.py antes.xlsx antes.json depois.xlsx depois.json
Código de saída 0 = idênticos; 1 = há diferenças (listadas).
"""
import json
import re
import sys
import zipfile

from openpyxl import load_workbook

HORA = re.compile(r"Atualizado em \d\d/\d\d/\d{4} \d\d:\d\d")


def norm(v):
    return HORA.sub("Atualizado em X", v) if isinstance(v, str) else v


def estilo(c):
    return (c.number_format, repr(c.font), repr(c.fill), repr(c.border), repr(c.alignment))


def livro(caminho):
    wb = load_workbook(caminho)
    out = {"folhas": wb.sheetnames, "calc": repr(wb.calculation.fullCalcOnLoad)}
    for ws in wb.worksheets:
        celulas = {}
        for linha in ws.iter_rows():
            for c in linha:
                if c.value is not None or c.has_style:
                    celulas[c.coordinate] = (norm(c.value), estilo(c))
        out[ws.title] = {
            "celulas": celulas,
            "merges": sorted(str(m) for m in ws.merged_cells.ranges),
            "larguras": {k: (v.width, v.hidden) for k, v in ws.column_dimensions.items()},
            "alturas": {k: v.height for k, v in ws.row_dimensions.items() if v.height is not None},
            "freeze": ws.freeze_panes, "grelha": ws.sheet_view.showGridLines, "tab": repr(ws.sheet_properties.tabColor),
            "tabelas": sorted((t.displayName, t.ref, t.tableStyleInfo.name) for t in ws.tables.values()),
            "validacoes": sorted(repr((d.type, d.formula1, d.formula2, str(d.sqref), d.error))
                                 for d in ws.data_validations.dataValidation),
            "condicional": sorted(repr((str(r.sqref), [(x.formula, repr(x.dxf)) for x in r.rules]))
                                  for r in ws.conditional_formatting),
        }
    with zipfile.ZipFile(caminho) as z:
        out["xml"] = {n: z.read(n).decode("utf-8") for n in sorted(z.namelist())
                      if re.match(r"xl/(charts/chart\d+|drawings/drawing\d+)\.xml$", n)}
        for n in z.namelist():   # ignoredErrors acrescentado ao XML
            if re.match(r"xl/worksheets/sheet\d+\.xml$", n):
                out.setdefault("ignored", {})[n] = re.findall(r"<ignoredErrors>.*?</ignoredErrors>", z.read(n).decode())
    return out


def dados(caminho):
    d = json.loads(open(caminho, encoding="utf-8").read())
    for k in ("gerado_em_utc", "sha256_excel", "gerador", "ficheiro_excel"):
        d["metadados"][k] = "X"
    return d


def diferencas(a, b, caminho=""):
    if type(a) is not type(b):
        return [f"{caminho}: tipo {type(a).__name__} != {type(b).__name__}"]
    if isinstance(a, dict):
        out = []
        for k in sorted(set(a) | set(b), key=str):
            if k not in a or k not in b:
                out.append(f"{caminho}/{k}: só num dos lados")
            else:
                out += diferencas(a[k], b[k], f"{caminho}/{k}")
        if list(a) != list(b) and not out:
            out.append(f"{caminho}: ordem das chaves diferente")
        return out
    if isinstance(a, (list, tuple)):
        if len(a) != len(b):
            return [f"{caminho}: comprimento {len(a)} != {len(b)}"]
        out = []
        for i, (x, y) in enumerate(zip(a, b, strict=True)):
            out += diferencas(x, y, f"{caminho}[{i}]")
        return out
    return [] if a == b else [f"{caminho}: {a!r} != {b!r}"[:400]]


if __name__ == "__main__":
    ref_x, ref_j, novo_x, novo_j = sys.argv[1:5]
    d = diferencas(livro(ref_x), livro(novo_x)) + diferencas(dados(ref_j), dados(novo_j))
    print(f"DIFERENÇAS: {len(d)}")
    for linha in d[:40]:
        print("  " + linha)
    sys.exit(1 if d else 0)
