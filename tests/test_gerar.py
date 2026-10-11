import hashlib
import json
import re
import zipfile
from argparse import Namespace
from datetime import date, datetime, timedelta
from pathlib import Path

import jsonschema
import pytest
from openpyxl import load_workbook

from zepp_transformacao import gerar, sintetico, verificacao
from zepp_transformacao.medidas import ler_existentes

ESQUEMA = json.loads((Path(__file__).resolve().parents[1] / "schema" / "dados-zepp-life-ia.schema.json").read_text(encoding="utf-8"))
FOLHAS = ["Painel", "Resumo Mensal", "Resumo Semanal", "Medidas Manuais", "Medições Corporais", "Critérios e Qualidade"]


def correr(projeto, existente=None, hoje="2026-05-02", nome="livro", excel_real=False):
    pasta = projeto["pasta"]
    args = Namespace(origem=str(projeto["origem"]), excel=str(pasta / f"{nome}.xlsx"), json=str(pasta / f"{nome}.json"),
                     existente=str(existente) if existente else None, config=str(projeto["config"]), hoje=hoje,
                     sem_excel_real=not excel_real, graficos=None)
    resultado = gerar.executar(args)
    return Path(args.excel), json.loads(Path(args.json).read_text(encoding="utf-8")), resultado


def test_json_valido_reconciliado_e_sem_segredos(projeto):
    livro, dados, _ = correr(projeto)
    jsonschema.validate(dados, ESQUEMA)
    assert dados["metadados"]["sha256_excel"] == hashlib.sha256(livro.read_bytes()).hexdigest()
    c = dados["criterios_selecao"]["contagens"]
    assert c["aceites"] == len(dados["medicoes_corporais"])
    assert c["incompletas"] == len(dados["pesagens_excluidas"])
    assert "AABBCCDDEEFF0011" not in json.dumps(dados)          # identificador do dispositivo

    def chaves(valor):
        if isinstance(valor, dict):
            for k, v in valor.items():
                yield k
                yield from chaves(v)
        elif isinstance(valor, list):
            for v in valor:
                yield from chaves(v)
    proibidas = {"token", "Token", "ProfileID", "Password", "Account", "email", "EncryptImpedance", "Source", "User"}
    assert not proibidas & set(chaves(dados))
    ids = [m["id_medicao"] for m in dados["medicoes_corporais"]]
    assert ids == sorted(ids) and len(set(ids)) == len(ids)
    assert sum(r["leituras"] for r in dados["resumos_semanais"]) == len(ids)
    assert sum(r["leituras"] for r in dados["resumos_mensais"]) == len(ids)
    semanas = {m["semana_inicio"] for m in dados["medidas_manuais"]}
    assert semanas == {r["inicio"] for r in dados["resumos_semanais"]}


def test_ausencias_sao_null_e_nunca_zero(projeto):
    _, dados, _ = correr(projeto)
    ferias = [m for m in dados["medidas_manuais"] if m["semana_inicio"] in ("2026-02-15", "2026-02-22")]
    assert ferias and all(m["cintura_cm"] is None and m["estado"] == "por_preencher" for m in ferias)
    esquecida = next(m for m in dados["medidas_manuais"] if m["semana_inicio"] == "2026-02-01")
    assert esquecida["estado"] == "sem_medicao" and esquecida["data_medicao"] is None
    medida = next(m for m in dados["medidas_manuais"] if m["cintura_cm"] is not None)
    assert len(medida["cintura_leituras_cm"]) == 3


def test_livro_estrutura_e_graficos(projeto):
    livro, dados, resultado = correr(projeto)
    wb = load_workbook(livro)
    assert wb.sheetnames == FOLHAS
    assert all(ws.sheet_state == "visible" for ws in wb.worksheets)
    assert wb["Medidas Manuais"]["A3"].value == "Semana"
    assert wb["Medições Corporais"].column_dimensions["Q"].hidden
    lacunas = 0
    semanas = resultado["semanas"]
    medidos = [i for i, s in enumerate(semanas) if s.cintura is not None]
    lacunas = sum(1 for a, b in zip(medidos, medidos[1:]) if b - a > 1)
    with zipfile.ZipFile(livro) as z:
        graficos = [z.read(n).decode("utf-8") for n in sorted(z.namelist()) if re.match(r"xl/charts/chart\d+\.xml", n)]
    assert len(graficos) == 5
    manuais = [g for g in graficos if "'Medidas Manuais'" in g]
    assert len(manuais) == 2 and all('dispBlanksAs val="span"' in g for g in manuais)
    cintura = next(g for g in manuais if "$C$4" in g)
    assert lacunas >= 1 and cintura.count('prstDash val="dash"') == lacunas
    assert all("<c:smooth val=\"1\"" not in g for g in graficos)


def test_preserva_edicoes_manuais_e_estende_periodos(projeto):
    livro, _, _ = correr(projeto, nome="v1")
    wb = load_workbook(livro)
    ws = wb["Medidas Manuais"]
    # semana 2 (11/01): corrigir data e cintura, apagar pescoço e nota
    ws["B5"], ws["C5"], ws["E5"], ws["G5"] = date(2026, 1, 13), 93.3, None, None
    # semana 7 (15/02, férias): registar a ausência
    ws["G10"] = "Sem medição: férias"
    wb.save(livro)
    antes = ler_existentes(livro)

    exportacao = json.loads(projeto["origem"].read_text(encoding="utf-8"))
    ultima = max(datetime.fromisoformat(r["Date"]) for r in exportacao)
    for d in range(1, 40):
        exportacao.append(sintetico.registo(ultima + timedelta(days=d), 75, 19.5, 58.8))
    projeto["origem"].write_text(json.dumps(exportacao), encoding="utf-8")
    livro2, dados, _ = correr(projeto, existente=livro, hoje=(ultima + timedelta(days=39)).date().isoformat(), nome="v2")
    depois = ler_existentes(livro2)
    for inicio, semana in antes.items():
        d = depois[inicio]
        assert (d.data, d.cintura, d.pescoco, d.notas) == (semana.data, semana.cintura, semana.pescoco, semana.notas)
    s = depois[date(2026, 1, 11)]
    assert (s.data, s.cintura, s.pescoco, s.notas) == (date(2026, 1, 13), 93.3, None, None)
    assert depois[date(2026, 2, 15)].notas == "Sem medição: férias"
    assert len(depois) > len(antes)
    assert dados["resumos_mensais"][-1]["rotulo"] == (ultima + timedelta(days=39)).strftime("%Y-%m")
    assert dados["medicoes_corporais"][-1]["peso_kg"] == 75


def test_erro_quando_o_livro_existente_esta_corrompido(projeto, tmp_path):
    estragado = tmp_path / "estragado.xlsx"
    estragado.write_bytes(b"isto nao e um xlsx")
    with pytest.raises(RuntimeError, match="preservar as medidas manuais"):
        correr(projeto, existente=estragado)


def test_validado_no_excel_real(projeto, tmp_path):
    livro, dados, resultado = correr(projeto, nome="real")
    try:
        valores = verificacao.calcular_no_excel(livro, tmp_path / "graficos")
    except verificacao.ExcelIndisponivel:
        pytest.skip("Microsoft Excel não disponível")
    from zepp_transformacao import resumos
    selecao, semanas = resultado["selecao"], resultado["semanas"]
    leituras = selecao.aceites
    diferencas = verificacao.comparar(
        valores, resultado["semanais"], resultado["mensais"], resumos.medias_moveis(leituras),
        [l.massa_gorda for l in leituras], resumos.indicadores(leituras, semanas),
        [resumos.variacoes_cintura(semanas)[s.inicio] for s in semanas], [s.estado for s in semanas])
    assert diferencas == []
    assert len(list((tmp_path / "graficos").glob("*.png"))) == 5


def test_ultima_pesagem_na_mudanca_de_hora(projeto):
    # Duas pesagens às 01:30 em Lisboa (UTC+1 e depois UTC+0): a última é a de maior instante UTC.
    from datetime import timezone
    exportacao = json.loads(projeto["origem"].read_text(encoding="utf-8"))
    exportacao += [sintetico.registo(datetime(2026, 10, 25, 0, 30, tzinfo=timezone.utc), 80, 20, 58),
                   sintetico.registo(datetime(2026, 10, 25, 1, 30, tzinfo=timezone.utc), 79, 20, 58)]
    projeto["origem"].write_text(json.dumps(exportacao), encoding="utf-8")
    livro, dados, _ = correr(projeto, hoje="2026-10-26", excel_real=True)   # compara com o Excel real, se existir
    assert dados["medicoes_corporais"][-1]["peso_kg"] == 79
    assert dados["indicadores"]["peso_kg"]["ultima_leitura"] == 79
    assert dados["metadados"]["ultima_medicao_corporal_utc"] == "2026-10-25T01:30:00Z"
    formula = load_workbook(livro)["Painel"]["B5"].value
    assert "$R$" in formula and "MATCH(MAX(" in formula      # procura pelo instante UTC, não pela hora local


def test_gera_sem_nenhuma_data_manual(projeto):
    livro, _, _ = correr(projeto, nome="v1")
    wb = load_workbook(livro)
    ws = wb["Medidas Manuais"]
    for linha in range(4, ws.max_row + 1):
        for coluna in "BCEG":
            ws[f"{coluna}{linha}"] = None
    wb.save(livro)
    _, dados, _ = correr(projeto, existente=livro, nome="v2", excel_real=True)
    assert all(m["data_medicao"] is None and m["estado"] == "por_preencher" for m in dados["medidas_manuais"])
    assert dados["indicadores"]["cintura_cm"] is None and dados["indicadores"]["pescoco_cm"] is None


def test_comparacao_acusa_folhas_em_falta():
    diferencas = verificacao.comparar({"folhas": {}, "erros": []}, [], [], [], [], {}, [], [])
    assert diferencas and "Folhas não lidas" in diferencas[0]
